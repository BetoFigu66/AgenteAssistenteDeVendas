import argparse
import re
import sys
from copy import deepcopy
from datetime import datetime
from pathlib import Path

import yaml
from gera_grafico_cobertura_ponderada import _saida_para_sprint, _ultima_sprint_numero, gerar_grafico
from pptx import Presentation
from pptx.oxml.ns import qn

# Reconfigurar stdout/stderr para UTF-8 com fallback. No Windows, o console
# padrao (PowerShell/cmd) usa cp1252 e prints de DEBUG com caracteres
# Unicode (emojis, pontuacao tipografica, U+FFFD) derrubam o script ANTES
# de prs.save(...), deixando o PPTX antigo corrompido no disco.
for _stream in (sys.stdout, sys.stderr):
    try:
        _stream.reconfigure(encoding="utf-8", errors="replace")
    except (AttributeError, ValueError):
        pass

PROJECT_ROOT = Path(__file__).resolve().parent
DEFAULT_TEMPLATE = PROJECT_ROOT / "artefatos" / "gerente_de_projetos" / "sprint_review_template_v01.pptx"
EMOJI_PRIO = {"alta": "🔴", "media": "🟡", "baixa": "🟢"}

# Diretorio para PNGs de graficos gerados sob demanda durante a montagem do PPTX.
GRAFICOS_DIR = Path(__file__).resolve().parents[2] / "artefatos" / "gerente_de_projetos" / "reports"

GRAFICO_TOKEN_RE = re.compile(r"\{\{GRAFICO:(\w+)\}\}")


def _gerar_grafico_png(nome: str, sprint_numero: int | None) -> Path:
    """Gera o PNG do grafico identificado por `nome` para a sprint informada.

    Dispatch atualmente conhece:
      - cobertura_ponderada: barras empilhadas (realizado/faltante por REQ).
    """
    if nome == "cobertura_ponderada":
        numero = sprint_numero if sprint_numero is not None else _ultima_sprint_numero()
        saida = _saida_para_sprint(numero)
        gerar_grafico(saida, sprint_numero=numero)
        return saida
    raise ValueError(f"Grafico desconhecido: {nome}")


def _processar_graficos(prs, tokens_usados: set, sprint_numero: int | None = None):
    """Substitui shapes cujo texto contem {{GRAFICO:nome}} por uma imagem PNG.

    A imagem ocupa a mesma posicao/tamanho do shape original, que e removido.
    O PNG e gerado on-demand via _gerar_grafico_png, usando sprint_numero
    (extraido do YAML do report) quando aplicavel.
    """
    for slide in prs.slides:
        # snapshot da lista de shapes para permitir remocao durante iteracao
        for shape in list(slide.shapes):
            if not shape.has_text_frame:
                continue
            texto = "\n".join("".join(r.text for r in p.runs) for p in shape.text_frame.paragraphs)
            m = GRAFICO_TOKEN_RE.search(texto)
            if not m:
                continue
            nome = m.group(1)
            try:
                png_path = _gerar_grafico_png(nome, sprint_numero)
            except Exception as e:
                print(f"[ERRO] Falha ao gerar grafico '{nome}': {e}", file=sys.stderr)
                continue
            left, top, width, height = shape.left, shape.top, shape.width, shape.height
            sp = shape._element
            sp.getparent().remove(sp)
            slide.shapes.add_picture(str(png_path), left, top, width=width, height=height)
            tokens_usados.add(f"GRAFICO:{nome}")


def _flatten_dict(d: dict, parent_key: str = "") -> dict:
    result = {}
    for key, value in d.items():
        new_key = f"{parent_key}.{key}" if parent_key else key
        if isinstance(value, dict):
            result.update(_flatten_dict(value, new_key))
        elif not isinstance(value, list):
            result[new_key] = _to_scalar_string(value)
    return result


def _fmt_data(s: str) -> str:
    try:
        return datetime.strptime(s, "%Y-%m-%d").strftime("%d/%m/%Y")
    except Exception:
        return s or ""


def _is_iso_date(value):
    return isinstance(value, str) and re.fullmatch(r"\d{4}-\d{2}-\d{2}", value) is not None


def _to_scalar_string(value):
    if _is_iso_date(value):
        return _fmt_data(value)
    if isinstance(value, list):
        # Lista vira "a, b, c" em vez de "['a', 'b', 'c']" (colchetes e aspas
        # ficam feios quando renderizados em texto no PPT).
        return ", ".join(_to_scalar_string(v) for v in value)
    if isinstance(value, str):
        # YAML folded scalars (`>`) preservam newline final por padrao, o que
        # vira linha em branco quando concatenado a outro texto no template.
        # Strip de bordas resolve sem afetar conteudo interno.
        return value.strip()
    return str(value or "")


ALIASED_KEYS = {
    "backlogpendente": "backlog_pendente",
    "proximasprint": "proximo_sprint",
    "backlog": "backlogpendente",
    # Renomeacao: bloqueios -> riscos_e_impedimentos (estrutura agora e
    # lista de dicts com {nome, tipo, motivo, acao_esperada}, nao mais
    # lista de strings). O alias mantem compat com templates antigos.
    "bloqueios": "riscos_e_impedimentos",
}
ALIASED_KEYS.update({v: k for k, v in list(ALIASED_KEYS.items())})


def _token_key_variants(chave: str):
    variants = {chave}
    if chave in ALIASED_KEYS:
        variants.add(ALIASED_KEYS[chave])
    if "_" in chave:
        variants.add(chave.replace("_", ""))
    return variants


def _wrap_token(name: str) -> str:
    return f"{{{{{name}}}}}"


def _parse_placeholder(token_name: str):
    token_name = token_name.strip()
    for sep in (":", "|"):
        if sep in token_name:
            base, suffix = token_name.rsplit(sep, 1)
            if suffix.isdigit():
                return base, {"limit": int(suffix)}
    return token_name, {}


def _chunk_list(values: list, limit: int):
    return [values[i : i + limit] for i in range(0, len(values), limit)]


def _extract_placeholders(text: str):
    return re.findall(r"\{\{\s*([^}\s]+)\s*\}\}", text)


def _split_token_path(token: str, dados: dict):
    token = token.strip()
    if token in dados or _resolve_yaml_key(token, dados) in dados:
        return [token]
    if "." in token:
        return token.split(".")
    if "_" in token:
        parts = token.split("_")
        for split in range(1, len(parts)):
            base = "_".join(parts[:split])
            if _resolve_yaml_key(base, dados) in dados:
                return [base] + parts[split:]
    return [token]


def _resolve_yaml_key(chave: str, dados: dict):
    if chave in dados:
        return chave
    for alias in _token_key_variants(chave):
        if alias in dados:
            return alias
    return chave


def _resolve_yaml_path_value(path_parts: list, dados: dict):
    chave = _resolve_yaml_key(path_parts[0], dados)
    if chave not in dados:
        return None
    valor = dados[chave]
    remaining = path_parts[1:]

    def _descend(current, parts):
        if not parts:
            if isinstance(current, list):
                if all(isinstance(item, dict) for item in current):
                    return [_format_list_item(chave, item) for item in current]
                return [str(item) for item in current]
            if isinstance(current, dict):
                # Dict no final do path: vira lista de bullets (uma linha por
                # chave). Caso de uso: {{cobertura_sprint.por_req}}.
                if current:
                    return [_formatar_cobertura_req_item(k, v) for k, v in current.items()]
                return "(vazio)"
            return _to_scalar_string(current)

        part, rest = parts[0], parts[1:]
        if isinstance(current, dict):
            if part in current:
                return _descend(current[part], rest)
            if part in ALIASED_KEYS and ALIASED_KEYS[part] in current:
                return _descend(current[ALIASED_KEYS[part]], rest)
            return _to_scalar_string("") if not rest else ""

        if isinstance(current, list):
            results = []
            for item in current:
                if isinstance(item, dict):
                    if part in item:
                        value = item[part]
                    elif part in ALIASED_KEYS and ALIASED_KEYS[part] in item:
                        value = item[ALIASED_KEYS[part]]
                    else:
                        value = ""
                    sub = _descend(value, rest)
                    if isinstance(sub, list):
                        results.extend(sub)
                    else:
                        results.append(sub)
                else:
                    results.append(_to_scalar_string("") if rest else _to_scalar_string(item))
            return results

        return _to_scalar_string(current)

    return _descend(valor, remaining)


def _collect_slide_limit_placeholders(slide, dados):
    found = []

    def scan_text(texto):
        for token_name in _extract_placeholders(texto):
            base, opts = _parse_placeholder(token_name)
            if "limit" not in opts:
                continue
            resolved = _resolve_yaml_path_value(_split_token_path(base, dados), dados)
            if isinstance(resolved, list) and len(resolved) > opts["limit"]:
                found.append((token_name, base, opts, resolved))

    for shape in slide.shapes:
        if shape.has_text_frame:
            for paragraph in shape.text_frame.paragraphs:
                scan_text("".join(run.text for run in paragraph.runs))
        if shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:
                    for paragraph in cell.text_frame.paragraphs:
                        scan_text("".join(run.text for run in paragraph.runs))
    return found


def _process_slide_list_limits(prs, dados):
    replacements = {}
    original_slides = list(prs.slides)
    sldIdLst = prs.slides._sldIdLst

    # Processar em ordem reversa: clones de iteracoes posteriores ja estarao
    # posicionados quando processarmos um slide de indice menor, garantindo
    # que all_sldIds[idx] continue apontando para o template correto. Sem
    # isso, os clones inseridos em iteracoes anteriores deslocam os indices
    # e o template original (chunk[0]) acaba sendo movido para o fim do deck.
    for idx, slide in reversed(list(enumerate(original_slides))):
        placeholders = _collect_slide_limit_placeholders(slide, dados)
        if not placeholders:
            continue

        raw_token, base, opts, resolved = placeholders[0]
        chunks = _chunk_list(resolved, opts["limit"])
        if len(chunks) <= 1:
            continue

        clones = []
        for _ in chunks[1:]:
            clones.append(_clonar_slide(prs, slide))

        all_sldIds = sldIdLst.findall(qn("p:sldId"))
        template_sldId = all_sldIds[idx]
        new_sldIds = all_sldIds[-len(clones) :]
        for sldId in reversed(new_sldIds):
            template_sldId.addnext(sldId)

        chunk_token = _wrap_token(raw_token)
        for slide_instance, chunk in zip([slide] + clones, chunks):
            replacements[id(slide_instance)] = {
                chunk_token: [_format_list_item(base, item) if isinstance(item, dict) else str(item) for item in chunk]
            }

    return replacements


def _describe_structure(value):
    if isinstance(value, dict):
        return {
            "type": "dict",
            "fields": list(value.keys()),
            "children": {k: _describe_structure(v) for k, v in value.items()},
            "repeat": False,
        }
    if isinstance(value, list):
        item_meta = _describe_structure(value[0]) if value else {"type": "scalar", "repeat": False}
        return {
            "type": "list",
            "item": item_meta,
            "item_type": item_meta["type"],
            "count": len(value),
            "repeat": True,
        }
    return {"type": "scalar", "repeat": False}


def _build_yaml_metadata(dados: dict) -> dict:
    return {key: _describe_structure(value) for key, value in dados.items()}


def _format_list_item(chave: str, item):
    if isinstance(item, str):
        return item
    if isinstance(item, dict):
        # Comentado para evitar montagem de strings legíveis específicas
        # if "titulo" in item and "agente" in item:
        #     desc = item.get("descricao", "")
        #     texto = f"{item.get('titulo', '')} — {item.get('agente', '')}"
        #     return f"{texto}: {desc}" if desc else texto
        # if "titulo" in item and "prioridade" in item:
        #     return f"{EMOJI_PRIO.get(item.get('prioridade', 'media'), '⚪')}
        # {item.get('titulo', '')} ({item.get('prioridade', 'media')})"
        # prioritized = [item.get(field) for field in (
        # "titulo", "nome", "descricao", "agente", "prioridade") if field in item]
        # if prioritized:
        #     return " — ".join(str(v) for v in prioritized if v)
        return str(item)
    return str(item)


def _preparar_tokens_pptx(dados: dict, metadata: dict):
    flat_dados = _flatten_dict(dados)
    tokens_simples = {f"{{{{{k}}}}}": v for k, v in flat_dados.items()}

    tokens_lista = {}
    for chave, meta in metadata.items():
        if meta["type"] != "list":
            continue

        itens = dados.get(chave) or []
        valores = (
            [_format_list_item(chave, item) for item in itens]
            if meta["item_type"] == "dict"
            else [str(item) for item in itens]
        )
        valores = valores or ["(nenhum)"]

        for token_key in _token_key_variants(chave):
            tokens_lista[_wrap_token(token_key)] = valores

        if itens and meta["item_type"] == "dict":
            primeiro_item = itens[0]
            for field in meta["item"]["fields"]:
                valor = primeiro_item.get(field, "")
                for token_key in _token_key_variants(chave):
                    tokens_simples[_wrap_token(f"{token_key}.{field}")] = _to_scalar_string(valor)
                    tokens_simples[_wrap_token(f"{token_key}_{field}")] = _to_scalar_string(valor)

    return tokens_simples, tokens_lista


def _substituir_texto_em_p_xml(p_element, alvo: str, novo: str):
    """Substitui ocorrencias do token preservando a formatacao por run.

    O token frequentemente cruza varios runs (ex.: "{{", "nome", "}}"), e o
    parágrafo pode ter runs com formatacao diferente antes/depois do token
    (ex.: "Severidade" em bold antes de ": {{token}}" em fonte normal).

    Algoritmo: para cada ocorrencia, localiza o run/offset do inicio e fim
    do token. O texto substituto fica no run que contem o inicio (herdando
    seu estilo), runs intermediarios sao zerados, e o restante do run final
    (o que vier depois de "}}") e mantido. Assim, runs antes do inicio do
    token preservam integralmente texto e estilo.
    """
    runs = p_element.findall(qn("a:r"))
    if not runs:
        return
    a_t_elements = []
    textos = []
    for r in runs:
        t = r.find(qn("a:t"))
        a_t_elements.append(t)
        textos.append(t.text or "" if t is not None else "")

    while True:
        total = "".join(textos)
        idx_start = total.find(alvo)
        if idx_start < 0:
            break
        idx_end = idx_start + len(alvo)

        run_start = run_end = None
        offset_start = offset_end = 0
        acc = 0
        for i, t in enumerate(textos):
            nxt = acc + len(t)
            if run_start is None and idx_start < nxt:
                run_start, offset_start = i, idx_start - acc
            if idx_end <= nxt:
                run_end, offset_end = i, idx_end - acc
                break
            acc = nxt
        if run_start is None or run_end is None:
            break

        if run_start == run_end:
            r = textos[run_start]
            textos[run_start] = r[:offset_start] + novo + r[offset_end:]
        else:
            textos[run_start] = textos[run_start][:offset_start] + novo
            for i in range(run_start + 1, run_end):
                textos[i] = ""
            textos[run_end] = textos[run_end][offset_end:]

    for t_elem, novo_t in zip(a_t_elements, textos):
        if t_elem is not None:
            t_elem.text = novo_t


def _substituir_texto_no_paragrafo(paragraph, alvo: str, novo: str):
    _substituir_texto_em_p_xml(paragraph._p, alvo, novo)


def _replace_tokens_in_paragraph(paragraph, replacements, tokens_usados):
    texto_total = "".join(run.text for run in paragraph.runs)
    for token, valor in replacements.items():
        if token in texto_total:
            tokens_usados.add(token.replace("{{", "").replace("}}", ""))
            _substituir_texto_no_paragrafo(paragraph, token, valor)


def _replace_tokens_in_p_element(p_element, replacements):
    runs = p_element.findall(qn("a:r"))
    if not runs:
        return
    textos = []
    for r in runs:
        t = r.find(qn("a:t"))
        textos.append(t.text or "" if t is not None else "")
    texto_total = "".join(textos)
    for token, valor in replacements.items():
        if token in texto_total:
            _substituir_texto_em_p_xml(p_element, token, valor)


def _substituir_tokens_em_textframe(
    text_frame, tokens_simples, tokens_lista, tokens_usados, dados=None, slide_replacements=None
):
    slide_replacements = slide_replacements or {}
    paragrafos = list(text_frame.paragraphs)
    for paragraph in paragrafos:
        texto_paragrafo = "".join(run.text for run in paragraph.runs)
        if "{{" not in texto_paragrafo:
            continue

        placeholders = _extract_placeholders(texto_paragrafo)
        if not placeholders:
            continue

        replacements = {}
        list_tokens = []
        for token_name in placeholders:
            raw_token = _wrap_token(token_name)
            base_token, opts = _parse_placeholder(token_name)
            base_raw_token = _wrap_token(base_token)

            if raw_token in slide_replacements:
                valor = slide_replacements[raw_token]
            elif base_raw_token in tokens_simples:
                valor = tokens_simples[base_raw_token]
            elif base_raw_token in tokens_lista:
                valor = tokens_lista[base_raw_token]
            elif dados is not None:
                resolved = _resolve_yaml_path_value(_split_token_path(base_token, dados), dados)
                valor = resolved if resolved is not None else ""
            else:
                valor = ""

            # Dict resolvido vira lista de bullets (uma linha por chave).
            # Caso de uso principal: {{cobertura_sprint.por_req}} e similares,
            # onde cada entrada e { antes, depois, delta_pp }. Itens que nao
            # tem esse formato caem no fallback "chave: valor".
            if isinstance(valor, dict):
                if valor:
                    valor = [_formatar_cobertura_req_item(k, v) for k, v in valor.items()]
                else:
                    valor = "(vazio)"

            replacements[raw_token] = valor
            if isinstance(valor, list):
                list_tokens.append(raw_token)

        if not replacements:
            continue

        if not list_tokens:
            _replace_tokens_in_paragraph(paragraph, replacements, tokens_usados)
            continue

        primary_token = list_tokens[0]
        count = len(replacements[primary_token])
        original_p = deepcopy(paragraph._p)
        p_anchor = paragraph._p

        for index in range(count):
            item_replacements = {}
            for token, valor in replacements.items():
                if isinstance(valor, list):
                    if len(valor) == count:
                        item_replacements[token] = _to_scalar_string(valor[index])
                    else:
                        item_replacements[token] = _to_scalar_string(valor[0]) if valor else ""
                else:
                    item_replacements[token] = valor

            if index == 0:
                _replace_tokens_in_paragraph(paragraph, item_replacements, tokens_usados)
                continue

            novo_p = deepcopy(original_p)
            p_anchor.addnext(novo_p)
            p_anchor = novo_p
            _replace_tokens_in_p_element(novo_p, item_replacements)

    # DEBUG: Print tokens {{*}} that were not substituted
    for paragraph in text_frame.paragraphs:
        texto = "".join(run.text for run in paragraph.runs)
        tokens_nao_substituidos = re.findall(r"\{\{[^}]+\}\}", texto)
        for token in tokens_nao_substituidos:
            print(f"DEBUG: Token não substituído no template: {token}")


def _formatar_cobertura_req_item(req_name: str, info) -> str:
    """Formata uma entrada de cobertura_req como linha legivel.

    Exemplos:
      REQ-010: {antes: "~10%", depois: "~30%", delta_pp: "+20"}
        -> "REQ-010: antes ~10%, depois ~30%, delta +20"
      obs: "infra de suporte..."
        -> "obs: infra de suporte..."
    """
    if isinstance(info, dict):
        if {"antes", "depois", "delta_pp"} <= set(info.keys()):
            return f"{req_name}: antes {info['antes']}, depois {info['depois']}, delta {info['delta_pp']}"
        # fallback: junta os campos disponiveis
        partes = ", ".join(f"{k} {v}" for k, v in info.items())
        return f"{req_name}: {partes}"
    return f"{req_name}: {info}"


def _subtokens_por_item(chave: str, item, metadata: dict = None):
    """Retorna (tokens_simples, tokens_lista) para um item replicado por {{SLIDE:chave}}.

    Campos cujo valor e dict (ex.: feito.cobertura_req) sao convertidos em lista
    de linhas formatadas, indo para tokens_lista (renderizados como bullets pelo
    _substituir_tokens_em_textframe).
    """
    tokens_simples = {}
    tokens_lista = {}

    if isinstance(item, dict):
        for field, valor in item.items():
            if isinstance(valor, dict):
                # Dict aninhado: vira lista de bullets (uma linha por entrada).
                if field == "cobertura_req":
                    linhas = [_formatar_cobertura_req_item(k, v) for k, v in valor.items()]
                else:
                    linhas = [f"{k}: {_to_scalar_string(v)}" for k, v in valor.items()]
                if not linhas:
                    linhas = ["(vazio)"]
                for token_key in _token_key_variants(chave):
                    tokens_lista[_wrap_token(f"{token_key}.{field}")] = linhas
                    tokens_lista[_wrap_token(f"{token_key}_{field}")] = linhas
            else:
                for token_key in _token_key_variants(chave):
                    tokens_simples[_wrap_token(f"{token_key}.{field}")] = _to_scalar_string(valor)
                    tokens_simples[_wrap_token(f"{token_key}_{field}")] = _to_scalar_string(valor)
        for token_key in _token_key_variants(chave):
            tokens_simples[_wrap_token(f"{token_key}_texto")] = _format_list_item(chave, item)
            tokens_simples[_wrap_token(token_key)] = _format_list_item(chave, item)
        return tokens_simples, tokens_lista

    if isinstance(item, list):
        texto = ", ".join(str(v) for v in item)
        for token_key in _token_key_variants(chave):
            tokens_simples[_wrap_token(token_key)] = texto
            tokens_simples[_wrap_token(f"{token_key}_texto")] = texto
        return tokens_simples, tokens_lista

    for token_key in _token_key_variants(chave):
        tokens_simples[_wrap_token(token_key)] = _to_scalar_string(item)
        tokens_simples[_wrap_token(f"{token_key}_texto")] = _to_scalar_string(item)
    return tokens_simples, tokens_lista


def _ajustar_numeracao_sequencial(slide, start_at: int):
    """Injeta startAt em todos os <a:buAutoNum> dos paragrafos do slide.

    Por padrao, PPTX reinicia auto-numeracao em cada slide. Quando o slide
    e clonado N vezes (uma por item), todos comecam em "1.". Para numerar
    sequencialmente, basta setar startAt=N no buAutoNum de cada paragrafo
    numerado do clone N.
    """
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        for para in shape.text_frame.paragraphs:
            pPr = para._pPr
            if pPr is None:
                continue
            buAutoNum = pPr.find(qn("a:buAutoNum"))
            if buAutoNum is not None:
                buAutoNum.set("startAt", str(start_at))


def _clonar_slide(prs, slide):
    novo = prs.slides.add_slide(slide.slide_layout)
    spTree_novo = novo.shapes._spTree
    spTree_src = slide.shapes._spTree

    image_map = {}
    for rel_key, rel in slide.part.rels.items():
        if "image" in rel.reltype:
            image_map[rel_key] = rel.target_part

    for el in list(spTree_novo):
        if el.tag not in (qn("p:nvGrpSpPr"), qn("p:grpSpPr")):
            spTree_novo.remove(el)

    for shape_el in spTree_src:
        if shape_el.tag not in (qn("p:nvGrpSpPr"), qn("p:grpSpPr")):
            novo_shape_el = deepcopy(shape_el)
            blipFill = novo_shape_el.find(qn("p:blipFill"))
            if blipFill is not None:
                blip = blipFill.find(qn("a:blip"))
                if blip is not None:
                    old_rId = blip.rEmbed
                    if old_rId in image_map:
                        image_part = image_map[old_rId]
                        new_rId = novo.part.relate_to(
                            image_part,
                            "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image",
                        )
                        blip.rEmbed = new_rId
            spTree_novo.append(novo_shape_el)
    return novo


def _processar_slides_replicados(prs, dados: dict, tokens_usados: set = None, metadata: dict = None):
    if tokens_usados is None:
        tokens_usados = set()
    if metadata is None:
        metadata = _build_yaml_metadata(dados)

    SLIDE_TOKEN_RE = re.compile(r"\{\{SLIDE:(\w+)\}\}")

    def _detectar_chave(slide):
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    texto = "".join(r.text for r in para.runs)
                    m = SLIDE_TOKEN_RE.search(texto)
                    if m:
                        return m.group(1)
        return None

    slides_replicar = []
    for idx, slide in enumerate(prs.slides):
        chave = _detectar_chave(slide)
        if chave:
            chave_real = _resolve_yaml_key(chave, dados)
            slides_replicar.append((idx, chave, chave_real))
            tokens_usados.add(chave)

    sldIdLst = prs.slides._sldIdLst

    # Bug do python-pptx: PresentationPart._next_slide_partname usa
    # `len(sldIdLst) + 1` para gerar partnames de novos slides. Se removermos
    # sldIds do template DENTRO do loop, len(sldIdLst) diminui e o proximo
    # add_slide aloca um partname que ja esta em uso (colisao -> .pptx
    # corrompido com "Duplicate name" no zip). Solucao: adiar todas as
    # remocoes para o fim, depois que todos os clones ja foram criados.
    templates_para_remover = []

    for idx, chave, chave_real in reversed(slides_replicar):
        slide_template = prs.slides[idx]
        itens = dados.get(chave_real) or []
        if not itens:
            itens = (
                [f"(sem itens em '{chave}')"]
                if chave in ("bloqueios", "insights", "backlogpendente", "objetivos", "proximo_objetivos")
                else [{"titulo": f"(sem itens em '{chave}')"}]
            )

        novos_slides = []
        for item_idx, item in enumerate(itens, start=1):
            novo = _clonar_slide(prs, slide_template)
            sub_simples, sub_lista = _subtokens_por_item(chave, item, metadata)
            sub_simples[f"{{{{SLIDE:{chave}}}}}"] = ""
            for shape in novo.shapes:
                if shape.has_text_frame:
                    _substituir_tokens_em_textframe(shape.text_frame, sub_simples, sub_lista, set(), dados)
                if shape.has_table:
                    for row in shape.table.rows:
                        for cell in row.cells:
                            _substituir_tokens_em_textframe(cell.text_frame, sub_simples, sub_lista, set(), dados)
            # Numeracao sequencial entre clones: PPTX reinicia auto-num em
            # cada slide, entao injetamos startAt no <a:buAutoNum> de cada
            # paragrafo numerado para que o item N comece em N.
            _ajustar_numeracao_sequencial(novo, item_idx)
            novos_slides.append(novo)

        all_sldIds = sldIdLst.findall(qn("p:sldId"))
        template_sldId = all_sldIds[idx]
        n = len(novos_slides)
        novos_sldIds = all_sldIds[-n:]

        # addprevious(X) insere X imediatamente antes do template_sldId.
        # Iterando em ordem direta [c1, c2, c3], o resultado e:
        #   addprevious(c1) -> [c1, template]
        #   addprevious(c2) -> [c1, c2, template]
        #   addprevious(c3) -> [c1, c2, c3, template]
        # Iterar em reversed produziria a ordem invertida.
        for sldId in novos_sldIds:
            template_sldId.addprevious(sldId)

        templates_para_remover.append(template_sldId)

    # Remover sldIds dos templates somente APOS todos os clones terem sido
    # criados, garantindo que len(sldIdLst) nao decresca durante add_slide.
    for template_sldId in templates_para_remover:
        sldIdLst.remove(template_sldId)


def gerar_apresentacao_pptx(yaml_path, template_path=None, saida=None) -> Path:
    yaml_path = Path(yaml_path)
    if not yaml_path.exists():
        raise FileNotFoundError(f"Arquivo YAML nao encontrado: {yaml_path}")

    with open(yaml_path, "r", encoding="utf-8") as f:
        dados = yaml.safe_load(f)

    if template_path is None:
        template_path = DEFAULT_TEMPLATE
    template_path = Path(template_path)
    if not template_path.exists():
        raise FileNotFoundError(f"Template pptx nao encontrado: {template_path}")

    metadata = _build_yaml_metadata(dados)
    tokens_simples, tokens_lista = _preparar_tokens_pptx(dados, metadata)

    # DEBUG: Track all possible tokens from YAML
    todas_chaves_yaml = set()
    flat_dados = _flatten_dict(dados)
    todas_chaves_yaml.update(flat_dados.keys())
    todas_chaves_yaml.update(tokens_lista.keys())
    todas_chaves_yaml.update(token.replace("{{", "").replace("}}", "") for token in tokens_lista.keys())

    prs = Presentation(str(template_path))

    # DEBUG: Track used tokens
    tokens_usados = set()

    slide_replacements = _process_slide_list_limits(prs, dados)
    _processar_slides_replicados(prs, dados, tokens_usados, metadata)
    sprint_numero = dados.get("sprint_numero") or dados.get("numero_sprint") or dados.get("numero")
    try:
        sprint_numero = int(sprint_numero) if sprint_numero is not None else None
    except (TypeError, ValueError):
        sprint_numero = None
    _processar_graficos(prs, tokens_usados, sprint_numero=sprint_numero)

    for slide in prs.slides:
        slide_token_replacements = slide_replacements.get(id(slide), {})
        # print(f"DEBUG: Processando slide '{slide.shapes.title.text if slide.shapes.title else 'Sem título'}'")
        for shape in slide.shapes:
            if shape.has_text_frame:
                _substituir_tokens_em_textframe(
                    shape.text_frame, tokens_simples, tokens_lista,
                    tokens_usados, dados, slide_replacements=slide_token_replacements
                )
            if shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        _substituir_tokens_em_textframe(
                            cell.text_frame, tokens_simples, tokens_lista,
                            tokens_usados, dados, slide_replacements=slide_token_replacements
                        )
    # DEBUG: Print YAML keys that were not used
    chaves_nao_usadas = todas_chaves_yaml - tokens_usados
    if chaves_nao_usadas:
        print("DEBUG: Informações do YAML que não foram usadas:")
        for chave in sorted(chaves_nao_usadas):
            if chave in flat_dados:
                print(f"  - {chave}: {flat_dados[chave]}")
            elif chave in tokens_lista:
                print(f"  - {chave}: {tokens_lista[chave]}")

    if saida is None:
        saida = yaml_path.with_suffix(".pptx")
    saida = Path(saida)
    prs.save(str(saida))
    return saida


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Gerar PPTX de sprint a partir de um arquivo YAML e template PPTX.")
    parser.add_argument("yaml_file", help="Arquivo YAML com os dados da sprint")
    parser.add_argument(
        "template_file",
        nargs="?",
        default=str(DEFAULT_TEMPLATE),
        help="Template PPTX a ser usado (default: sprint_review_template_v01.pptx)",
    )
    parser.add_argument(
        "-o",
        "--saida",
        help="Arquivo de saída PPTX. Se omitido, o nome será o mesmo do YAML com extensão .pptx.",
    )

    args = parser.parse_args()
    yaml_file = args.yaml_file
    pptx_template = args.template_file
    saida = args.saida or Path(yaml_file).with_suffix(".pptx")

    saida = gerar_apresentacao_pptx(yaml_file, pptx_template, saida)
    print(f"Arquivo {saida} gerado baseado no arquivo de dados {yaml_file} e template {pptx_template}")
